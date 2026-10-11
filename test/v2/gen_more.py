"""Add more samples to ./samples: images, a rich DOCX, XLSX/CSV, Markdown and HTML."""

import csv
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageDraw, ImageFont

os.chdir("samples")
F = os.environ.get("FONT_BOLD", "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf")
im = Image.new("RGB", (900, 480), "white")
d = ImageDraw.Draw(im)
fnt = ImageFont.truetype(F, 28)
sm = ImageFont.truetype(F, 20)
d.text((30, 20), "분기별 매출 (단위: 백만원)", font=fnt, fill="#222")
for i, (q, v, c) in enumerate(
    [
        ("1분기", 120, "#3182f6"),
        ("2분기", 210, "#00a37a"),
        ("3분기", 180, "#f59e0b"),
        ("4분기", 300, "#ef4444"),
    ]
):
    x = 80 + i * 200
    d.rectangle([x, 440 - v, x + 120, 440], fill=c)
    d.text((x + 25, 445), q, font=sm, fill="#333")
    d.text((x + 35, 405 - v), str(v), font=sm, fill="#333")
im.save("chart.png")
im.convert("RGB").save("photo.jpg", quality=88)
doc = Document()
for s in ("Normal",):
    st = doc.styles[s]
    st.font.name = "맑은 고딕"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
sec = doc.sections[0]
sec.header.paragraphs[0].text = "꿀툴 · 사내 보고서 (대외비 아님)"
sec.footer.paragraphs[0].text = "— 1 —"
doc.add_heading("2026년 3분기 사업 보고서", 0)
p = doc.add_paragraph()
r = p.add_run("요약: ")
r.bold = True
r = p.add_run("3분기 매출은 전년 동기 대비 ")
r2 = p.add_run("23% 증가")
r2.bold = True
r2.font.color.rgb = RGBColor(0xEF, 0x44, 0x44)
p.add_run("했어요. 한글, English, 漢字, ①②③, ※ 기호가 모두 보여야 해요.")
doc.add_heading("1. 실적 표", level=1)
t = doc.add_table(rows=2, cols=4)
t.style = "Table Grid"
a = t.cell(0, 0).merge(t.cell(1, 0))
a.text = "구분"
b = t.cell(0, 1).merge(t.cell(0, 3))
b.text = "2026년 (단위: 백만원)"
for c, h in zip(t.rows[1].cells[1:], ["7월", "8월", "9월"]):
    c.text = h
for row in [
    ("매출", "58", "61", "61"),
    ("영업이익", "12", "15", "14"),
    ("신규 고객(명)", "1,204", "1,388", "1,512"),
]:
    cells = t.add_row().cells
    for c, v in zip(cells, row):
        c.text = v
doc.add_heading("2. 그림", level=1)
doc.add_picture("chart.png", width=Cm(14))
cap = doc.add_paragraph("그림 1. 분기별 매출")
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_heading("3. 다음 분기 계획", level=1)
for s in ["한글(HWP) → 워드 변환 출시", "PDF → 엑셀 표 추출 베타", "모바일 화면 개선"]:
    doc.add_paragraph(s, style="List Bullet")
doc.add_paragraph(
    "결론적으로, 대한민국 헌법 제1조 ① 대한민국은 민주공화국이다. ② 대한민국의 주권은 국민에게 있고, 모든 권력은 국민으로부터 나온다. "
    * 3
)
doc.save("report.docx")
wb = Workbook()
ws = wb.active
ws.title = "예산"
thin = Side(style="thin", color="999999")
bd = Border(left=thin, right=thin, top=thin, bottom=thin)
ws.merge_cells("A1:E1")
ws["A1"] = "2026년 부서별 예산 집행 현황"
ws["A1"].font = Font(bold=True, size=14)
ws["A1"].alignment = Alignment(horizontal="center")
hdr = ["부서", "예산(원)", "집행(원)", "잔액(원)", "집행률"]
ws.append([])
ws.append(hdr)
data = [
    ("기획팀", 12000000, 8450000),
    ("개발팀", 45000000, 39120000),
    ("디자인팀", 9000000, 4200000),
    ("마케팅팀", 30000000, 27750000),
    ("인사팀", 6000000, 2980000),
]
for i, (n, b_, s) in enumerate(data, start=4):
    ws.append([n, b_, s, f"=B{i}-C{i}", f"=C{i}/B{i}"])
ws.append(["합계", "=SUM(B4:B8)", "=SUM(C4:C8)", "=SUM(D4:D8)", "=C9/B9"])
for row in ws.iter_rows(min_row=3, max_row=9):
    for c in row:
        c.border = bd
        if c.column in (2, 3, 4):
            c.number_format = "#,##0"
        if c.column == 5:
            c.number_format = "0.0%"
for c in ws[3]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="3182F6")
    c.alignment = Alignment(horizontal="center")
for col, w in zip("ABCDE", (14, 16, 16, 16, 10)):
    ws.column_dimensions[col].width = w
wb.save("table.xlsx")
with open("table.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(hdr)
    for n, b_, s in data:
        w.writerow([n, b_, s, b_ - s, f"{s / b_:.1%}"])
open("notes.md", "w").write("""# 회의록 — 문서 변환 기능 확대

**일시**: 2026-10-08 (목) 15:00  ·  **참석**: 민형, 지수, 하늘

## 결정 사항
1. HWP ↔ DOCX 변환을 먼저 출시해요.
2. PDF → 엑셀은 *표가 있는 PDF*만 대상으로 해요.

## 표

| 항목 | 담당 | 기한 |
|---|---|---|
| 서버 이미지 | 민형 | 10/15 |
| 화면 | 지수 | 10/20 |

> 메모: 한셀(.cell)은 xlsx로 저장해서 올려 달라고 안내.

```js
console.log("코드 블록도 유지돼야 해요");
```
""")
open(
    "page.html", "w"
).write("""<!doctype html><html lang="ko"><meta charset="utf-8"><title>안내문</title>
<body style="font-family:'Noto Sans KR',sans-serif"><h1>꿀툴 이용 안내</h1>
<p>꿀툴은 <b>파일을 저장하지 않는</b> 무료 도구 모음이에요. <a href="https://kkultools.com">kkultools.com</a></p>
<h2>요금</h2><table border="1" cellpadding="6" style="border-collapse:collapse"><tr><th>구분</th><th>무료</th><th>Pro</th></tr>
<tr><td>문서 변환</td><td>-</td><td>월 300건</td></tr><tr><td>파일 크기</td><td>2GB(기기 안)</td><td>20MB(서버)</td></tr></table>
<ul><li>한글 HWP</li><li>워드 DOCX</li><li>엑셀 XLSX</li></ul><img src="chart.png" width="400"></body></html>""")
print("ok")
