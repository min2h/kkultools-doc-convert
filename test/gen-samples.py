#!/usr/bin/env python3
"""Generate Office sample files (Korean text, tables, lists, charts, sheets, slides).

Usage: gen-samples.py [OUT_DIR]
Needs python-docx, python-pptx and openpyxl. With soffice on PATH it also
writes doc/rtf/odt, ppt/odp and xls/ods copies.
"""

import os
import shutil
import subprocess
import sys

out = sys.argv[1] if len(sys.argv) > 1 else "samples"
os.makedirs(out, exist_ok=True)

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

d = Document()
st = d.styles["Normal"]
st.font.name = "맑은 고딕"
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
d.add_heading("2026년 꿀툴 문서 변환 시험 보고서", 0)
d.add_paragraph(
    "이 문서는 워드(DOCX) → PDF 변환 품질을 확인하려고 만든 시험 문서예요. 한글, English, 漢字, 숫자 1,234,567원, 특수문자 ①②③ ※ ★ 가 모두 제대로 보여야 해요."
)
d.add_heading("1. 표", level=1)
t = d.add_table(rows=1, cols=4)
t.style = "Table Grid"
for c, h in zip(t.rows[0].cells, ["구분", "1월", "2월", "합계"]):
    c.text = h
for row in [
    ("매출", "1,200만", "1,350만", "2,550만"),
    ("비용", "800만", "760만", "1,560만"),
    ("이익", "400만", "590만", "990만"),
]:
    cells = t.add_row().cells
    for c, v in zip(cells, row):
        c.text = v
d.add_heading("2. 목록", level=1)
for s in [
    "첫째, 파일은 저장하지 않아요.",
    "둘째, 변환이 끝나면 바로 지워요.",
    "셋째, 글꼴은 나눔·Noto로 대체해요.",
]:
    d.add_paragraph(s, style="List Number")
d.add_heading("3. 긴 본문", level=1)
for i in range(1, 40):
    d.add_paragraph(
        f"{i}번째 문단입니다. 대한민국 헌법 제1조 ① 대한민국은 민주공화국이다. ② 대한민국의 주권은 국민에게 있고, 모든 권력은 국민으로부터 나온다."
    )
d.save(os.path.join(out, "sample.docx"))

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

p = Presentation()
p.slide_width = Inches(13.333)
p.slide_height = Inches(7.5)
s = p.slides.add_slide(p.slide_layouts[0])
s.shapes.title.text = "꿀툴 발표 자료"
s.placeholders[1].text = "파워포인트(PPTX) → PDF 변환 시험 · 2026년 10월"
s = p.slides.add_slide(p.slide_layouts[1])
s.shapes.title.text = "핵심 요약"
tf = s.placeholders[1].text_frame
tf.text = "파일은 저장하지 않고 바로 지워요"
for line in [
    "변환 서버: Google Cloud (도쿄)",
    "한 달 300건, 하루 30건, 파일당 20MB",
    "한글(HWP)·워드·엑셀·PPT 지원",
]:
    tf.add_paragraph().text = line
s = p.slides.add_slide(p.slide_layouts[5])
s.shapes.title.text = "월별 변환 수"
cd = CategoryChartData()
cd.categories = ["7월", "8월", "9월", "10월"]
cd.add_series("변환 수", (120, 340, 560, 910))
s.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1.5), Inches(1.6), Inches(10), Inches(5.2), cd
)
p.save(os.path.join(out, "sample.pptx"))

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

wb = Workbook()
ws = wb.active
ws.title = "가계부"
ws.append(["날짜", "항목", "분류", "금액(원)"])
rows = [
    ("2026-10-01", "점심 김치찌개", "식비", 9000),
    ("2026-10-02", "지하철", "교통", 1550),
    ("2026-10-03", "월세", "주거", 550000),
    ("2026-10-04", "커피", "식비", 4500),
    ("2026-10-05", "책 구입", "문화", 18000),
]
for r in rows:
    ws.append(r)
ws.append(["", "", "합계", "=SUM(D2:D6)"])
for col, w in zip("ABCD", (14, 18, 10, 14)):
    ws.column_dimensions[col].width = w
ch = BarChart()
ch.title = "항목별 금액"
ch.add_data(Reference(ws, min_col=4, min_row=1, max_row=6), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=2, max_row=6))
ws.add_chart(ch, "F2")
ws2 = wb.create_sheet("명단")
ws2.append(["이름", "부서", "연락처"])
for i in range(1, 61):
    ws2.append([f"홍길동{i}", ["기획", "개발", "디자인"][i % 3], f"010-1234-{i:04d}"])
wb.save(os.path.join(out, "sample.xlsx"))

soffice = shutil.which("soffice")
if soffice:
    for src, fmts in (
        ("sample.docx", ("doc", "rtf", "odt")),
        ("sample.pptx", ("ppt", "odp")),
        ("sample.xlsx", ("xls", "ods")),
    ):
        for f in fmts:
            subprocess.run(
                [soffice, "--headless", "--convert-to", f, "--outdir", out, os.path.join(out, src)],
                check=True,
                capture_output=True,
            )
print("\n".join(sorted(os.listdir(out))))
